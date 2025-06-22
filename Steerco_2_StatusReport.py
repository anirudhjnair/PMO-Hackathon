import os
from pptx import Presentation
from openpyxl import Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

def extract_table_from_ppt_to_excel(ppt_file_path, excel_file_path):
    # Load the PowerPoint presentation
    presentation = Presentation(ppt_file_path)

    # Remove the Excel file if it already exists to ensure overwrite
    if os.path.exists(excel_file_path):
        os.remove(excel_file_path)

    # Create a new Excel workbook and select the active worksheet
    wb = Workbook()
    ws = wb.active
    ws.title = "PPT Table Data"

    row_counter = 1
    table_found = False

    # Iterate through slides
    for slide_index, slide in enumerate(presentation.slides, start=1):
        for shape in slide.shapes:
            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    ws.append(row_data)
                    row_counter += 1
                table_found = True
                break  # Only process the first table per slide

    # Add Excel table formatting if data was written
    if table_found:
        num_cols = len(table.columns)
        last_col_letter = chr(64 + num_cols) if num_cols <= 26 else 'Z'
        table_range = f"A1:{last_col_letter}{row_counter - 1}"
        excel_table = Table(displayName="PPTTable", ref=table_range)

        # Add a default style
        style = TableStyleInfo(name="TableStyleMedium9", showFirstColumn=False,
                               showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        excel_table.tableStyleInfo = style
        ws.add_table(excel_table)

    # Save the workbook
    wb.save(excel_file_path)
    print(f"Table data from PowerPoint saved to {excel_file_path}")

# File paths


def create_ppt_path(text):
    folder_path = r"C:\Users\A.JanardhananNair\OneDrive - Shell\Ani\Work\Learning\Hackathon\PMO"
    file_name = f"{text}_SteercoReport.pptx"
    full_path = os.path.join(folder_path, file_name)
    return full_path

def create_excel_path(text):
    folder_path = r"C:\Users\A.JanardhananNair\OneDrive - Shell\Ani\Work\Learning\Hackathon\PMO"
    file_name = f"{text}_StatusReport.xlsx"
    full_path = os.path.join(folder_path, file_name)
    return full_path

# Example usage
ppt_path = create_ppt_path("PRJ198847")
excel_path = create_excel_path("PRJ198847")

# Run the function
extract_table_from_ppt_to_excel(ppt_path, excel_path)
