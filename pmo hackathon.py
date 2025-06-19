from pptx import Presentation

# Load the presentation
#prs = Presentation("your_presentation.pptx")

prs = Presentation(r"C:\Users\A.JanardhananNair\OneDrive - Shell\Desktop\Test\Test PPT.pptx")

# Iterate through slides and extract text
for i, slide in enumerate(prs.slides):
    print(f"Slide {i+1}:")
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            print(shape.text)
